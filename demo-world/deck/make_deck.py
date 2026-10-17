"""Generate pitch.pptx: Whisk's half-finished 10-slide deck for the "Pitch in 5" demo.

Slides 3, 5, 6, 7 are deliberately "blank": placeholder text + empty image boxes.
Every shape has a stable name (shape.name) so the Deck agent can target it:
    s{N}_title, s{N}_body, s{N}_stat, s{N}_img_slot, s{N}_img_slot_2, ...

Run:  uv run --python 3.12 --with python-pptx python make_deck.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).with_name("pitch.pptx")

# Theme
BG = RGBColor(0x17, 0x20, 0x1A)       # deep ink green
CARD = RGBColor(0x22, 0x2E, 0x26)
MATCHA = RGBColor(0x8D, 0xBF, 0x5A)
CREAM = RGBColor(0xF6, 0xF1, 0xE7)
MUTED = RGBColor(0xA9, 0xB5, 0xAB)
TODO_RED = RGBColor(0xFF, 0x6B, 0x6B)
FONT = "Inter"  # falls back to Helvetica on machines without Inter

W, H = Inches(13.333), Inches(7.5)  # 16:9


def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def text(slide, name, left, top, width, height, value, size=20, color=CREAM,
         bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = value if isinstance(value, list) else [value]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        p.space_after = Pt(size * 0.4)
    return box


def kicker(slide, n, label):
    text(slide, f"s{n}_kicker", Inches(0.8), Inches(0.55), Inches(8), Inches(0.4),
         label.upper(), size=13, color=MATCHA, bold=True)


def page_no(slide, n):
    text(slide, f"s{n}_page", Inches(12.2), Inches(6.9), Inches(0.8), Inches(0.4),
         f"{n:02d}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def img_slot(slide, name, left, top, width, height, label="image"):
    """Empty dashed box: where the Deck agent drops a screenshot or the logo."""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.name = name
    shp.fill.background()
    shp.line.color.rgb = MUTED
    shp.line.width = Pt(1.5)
    shp.line.dash_style = MSO_LINE.DASH
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f"[ {label} ]"
    r.font.name = FONT
    r.font.size = Pt(14)
    r.font.color.rgb = MUTED
    return shp


def card(slide, name, left, top, width, height, title, body):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = CARD
    shp.line.fill.background()
    shp.adjustments[0] = 0.08
    text(slide, f"{name}_title", left + Inches(0.3), top + Inches(0.3), width - Inches(0.6), Inches(0.5),
         title, size=20, bold=True)
    text(slide, f"{name}_body", left + Inches(0.3), top + Inches(0.9), width - Inches(0.6), height - Inches(1.1),
         body, size=15, color=MUTED)


def new_slide(prs, n, label):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    set_bg(s)
    if label:
        kicker(s, n, label)
    page_no(s, n)
    return s


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1: Title (logo slot empty: Design agent fills it)
    s = new_slide(prs, 1, None)
    img_slot(s, "s1_img_slot", Inches(0.8), Inches(1.6), Inches(1.6), Inches(1.6), "logo")
    text(s, "s1_title", Inches(0.8), Inches(3.4), Inches(11), Inches(1.4), "Whisk", size=80, bold=True)
    text(s, "s1_body", Inches(0.8), Inches(4.8), Inches(11), Inches(0.8),
         "Matcha, delivered to your dorm.", size=28, color=MATCHA)
    text(s, "s1_footer", Inches(0.8), Inches(6.6), Inches(8), Inches(0.4),
         "LA Hacks 2026 · Seed pitch", size=13, color=MUTED)

    # 2: Problem (done)
    s = new_slide(prs, 2, "The problem")
    text(s, "s2_title", Inches(0.8), Inches(1.1), Inches(11.5), Inches(1.4),
         "Students spend $7 a day on matcha lattes.", size=44, bold=True)
    card(s, "s2_card1", Inches(0.8), Inches(3.2), Inches(3.7), Inches(2.8),
         "Expensive", "$150+/month at campus cafés for one drink a day.")
    card(s, "s2_card2", Inches(4.8), Inches(3.2), Inches(3.7), Inches(2.8),
         "Inconsistent", "Half the lattes are 80% syrup. Quality is a coin flip.")
    card(s, "s2_card3", Inches(8.8), Inches(3.2), Inches(3.7), Inches(2.8),
         "Inconvenient", "Lines between classes. No time during finals.")

    # 3: BLANK: Solution
    s = new_slide(prs, 3, "Our solution")
    text(s, "s3_title", Inches(0.8), Inches(1.1), Inches(11.5), Inches(1.2),
         "TODO: solution headline", size=44, bold=True, color=TODO_RED)
    text(s, "s3_body", Inches(0.8), Inches(2.6), Inches(6.2), Inches(3.6),
         ["TODO: 3 bullets on what Whisk does"], size=20, color=TODO_RED)
    img_slot(s, "s3_img_slot", Inches(7.4), Inches(2.4), Inches(5.1), Inches(4.0), "product photo / logo")

    # 4: How it works (done)
    s = new_slide(prs, 4, "How it works")
    text(s, "s4_title", Inches(0.8), Inches(1.1), Inches(11.5), Inches(1.2),
         "Three taps. Fresh tins every week.", size=44, bold=True)
    card(s, "s4_card1", Inches(0.8), Inches(3.0), Inches(3.7), Inches(3.0),
         "1 · Pick your vibe", "Ceremonial, latte-grade or flavored. Change anytime.")
    card(s, "s4_card2", Inches(4.8), Inches(3.0), Inches(3.7), Inches(3.0),
         "2 · We ship weekly", "Timed to your class schedule. Free shipping.")
    card(s, "s4_card3", Inches(8.8), Inches(3.0), Inches(3.7), Inches(3.0),
         "3 · Share & sip free", "Every referral earns 5 free credits.")

    # 5: BLANK: Product screenshots
    s = new_slide(prs, 5, "The product")
    text(s, "s5_title", Inches(0.8), Inches(1.1), Inches(11.5), Inches(1.0),
         "TODO: product screenshots", size=40, bold=True, color=TODO_RED)
    img_slot(s, "s5_img_slot", Inches(0.8), Inches(2.4), Inches(5.7), Inches(4.2), "screenshot: signup")
    img_slot(s, "s5_img_slot_2", Inches(6.9), Inches(2.4), Inches(5.6), Inches(4.2), "screenshot: landing page")

    # 6: BLANK: Market size
    s = new_slide(prs, 6, "Market")
    text(s, "s6_title", Inches(0.8), Inches(1.1), Inches(11.5), Inches(1.0),
         "TODO: market size", size=40, bold=True, color=TODO_RED)
    text(s, "s6_stat", Inches(0.8), Inches(2.5), Inches(6), Inches(2.0),
         "$?B", size=120, bold=True, color=TODO_RED)
    text(s, "s6_body", Inches(0.8), Inches(4.8), Inches(6.2), Inches(1.6),
         "TODO: source + one-line explanation", size=18, color=TODO_RED)
    img_slot(s, "s6_img_slot", Inches(7.4), Inches(2.4), Inches(5.1), Inches(4.0), "chart")

    # 7: BLANK: Pricing
    s = new_slide(prs, 7, "Business model")
    text(s, "s7_title", Inches(0.8), Inches(1.1), Inches(11.5), Inches(1.0),
         "TODO: pricing", size=40, bold=True, color=TODO_RED)
    text(s, "s7_stat", Inches(0.8), Inches(2.5), Inches(5.6), Inches(1.8),
         "$??/mo", size=96, bold=True, color=TODO_RED)
    text(s, "s7_body", Inches(0.8), Inches(4.5), Inches(5.8), Inches(2.0),
         "TODO: plans (standard vs student)", size=18, color=TODO_RED)
    img_slot(s, "s7_img_slot", Inches(7.0), Inches(2.4), Inches(5.5), Inches(4.0), "pricing table")

    # 8: Traction & roadmap (done)
    s = new_slide(prs, 8, "Traction & roadmap")
    text(s, "s8_title", Inches(0.8), Inches(1.1), Inches(11.5), Inches(1.2),
         "Launching at UCLA, then the UC system.", size=44, bold=True)
    card(s, "s8_card1", Inches(0.8), Inches(3.0), Inches(3.7), Inches(3.0),
         "Now", "Waitlist live. Supplier signed in Uji, Kyoto.")
    card(s, "s8_card2", Inches(4.8), Inches(3.0), Inches(3.7), Inches(3.0),
         "Q1", "UCLA pilot: 500 subscribers, dorm delivery.")
    card(s, "s8_card3", Inches(8.8), Inches(3.0), Inches(3.7), Inches(3.0),
         "Q3", "All 9 UC campuses + campus-café partnerships.")

    # 9: Team (done)
    s = new_slide(prs, 9, "Team")
    text(s, "s9_title", Inches(0.8), Inches(1.1), Inches(11.5), Inches(1.2),
         "One founder. A swarm of agents.", size=44, bold=True)
    text(s, "s9_body", Inches(0.8), Inches(2.6), Inches(11.5), Inches(3.5),
         ["Founder: CS @ UCLA, ex-barista, drinks too much matcha.",
          "Ops, design, web, deck & launch: Hive (yes, the agents built this deck)."],
         size=22, color=MUTED)

    # 10: Ask
    s = new_slide(prs, 10, "The ask")
    text(s, "s10_title", Inches(0.8), Inches(1.6), Inches(11.5), Inches(1.4),
         "Join the waitlist.", size=64, bold=True)
    text(s, "s10_body", Inches(0.8), Inches(3.2), Inches(7), Inches(1.0),
         "TODO: live URL", size=26, color=TODO_RED)
    img_slot(s, "s10_img_slot", Inches(9.3), Inches(1.6), Inches(3.2), Inches(3.2), "QR code")

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build()
    prs = Presentation(out)
    print(f"wrote {out} ({len(prs.slides)} slides)")
    for i, slide in enumerate(prs.slides, start=1):
        names = [sh.name for sh in slide.shapes]
        print(f"  slide {i}: {', '.join(names)}")
