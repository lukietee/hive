"""Deck worker: Claude edits a .pptx through small python-pptx steps (reusing DeckAgent's render +
Deck Viewer focus logic), each one re-rendered live in the Deck Viewer with the ghost cursor on the box."""

from __future__ import annotations

import asyncio
import time
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from ..deck_agent import DeckAgent
from ..orchestrator import llm
from .base import AgentSpec, DynamicAgent, WorkerResult

MAX_TURNS = 30


class DynDeck(DynamicAgent):
    """A DynamicAgent with DeckAgent's rendering/focus methods (DeckAgent hard-codes the 'deck' id)."""

    render = DeckAgent.render
    load = DeckAgent.load
    _box_frac = DeckAgent._box_frac
    _find_shape = staticmethod(DeckAgent.__dict__["_find_shape"].__func__)
    focus_rect = DeckAgent.focus_rect
    _focus = DeckAgent._focus
    set_text = DeckAgent.set_text

    def __init__(self, bus, spec: AgentSpec, pptx: Path, render_dir: Path,
                 base_url: str = "http://127.0.0.1:8766/render") -> None:
        super().__init__(bus, spec)
        self.pptx, self.render_dir, self.base_url = pptx, render_dir, base_url
        self.version = 0
        self.screen_mapper = None
        self.render_times: list[float] = []


def new_deck(path: Path, title: str) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    _add_slide(prs, title, "")
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def _add_slide(prs, title: str, body: str) -> int:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    n = len(prs.slides) - 1
    tb = s.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.3))
    tb.name = f"s{n}_title"
    tb.text_frame.text = title or " "
    r = tb.text_frame.paragraphs[0].runs
    if r:
        r[0].font.size, r[0].font.bold = Pt(40), True
    bb = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.6))
    bb.name = f"s{n}_body"
    bb.text_frame.word_wrap = True
    bb.text_frame.text = body or " "
    for p in bb.text_frame.paragraphs:
        for run in p.runs:
            run.font.size, run.font.color.rgb = Pt(24), RGBColor(0x33, 0x33, 0x38)
    return n


TOOLS: list[dict[str, Any]] = [
    {"name": "get_deck", "description": "List every slide with its shapes (name and current text).",
     "input_schema": {"type": "object", "properties": {}}},
    {"name": "set_text", "description": "Replace the text of one shape (by name, from get_deck). Use \\n for new lines.",
     "input_schema": {"type": "object", "properties": {"slide": {"type": "integer", "description": "0-based"},
                                                        "shape": {"type": "string"}, "text": {"type": "string"}},
                      "required": ["slide", "shape", "text"]}},
    {"name": "add_slide", "description": "Append a slide with a title and body text (\\n separates bullet lines).",
     "input_schema": {"type": "object", "properties": {"title": {"type": "string"}, "body": {"type": "string"}},
                      "required": ["title", "body"]}},
    {"name": "ask_human", "description": "Ask a decision with 2-4 options (best first); the human may answer freely.",
     "input_schema": {"type": "object", "properties": {"question": {"type": "string"},
                                                        "options": {"type": "array", "items": {"type": "string"}}},
                      "required": ["question", "options"]}},
    {"name": "finish", "description": "Call when the deck is done. summary = what you changed.",
     "input_schema": {"type": "object", "properties": {"summary": {"type": "string"}}, "required": ["summary"]}},
]

SYSTEM = """You are {name}, the deck agent on the Hive team. You edit a real PowerPoint deck in small steps; each step
is rendered live for an audience. Fill empty or TODO shapes with crisp, specific copy (short lines, no fluff);
add slides only when needed. Read the deck first with get_deck. Then call finish."""


class DeckWorker:
    def __init__(self, agent: DynDeck, answer_timeout: float = 30) -> None:
        self.agent = agent
        self.answer_timeout = answer_timeout
        self.human_notes: list[str] = []

    def _describe(self) -> str:
        prs = Presentation(self.agent.pptx)
        out = []
        for i, s in enumerate(prs.slides):
            out.append(f"Slide {i}:")
            for sh in s.shapes:
                text = sh.text_frame.text.replace("\n", " / ")[:160] if sh.has_text_frame else f"({sh.shape_type})"
                out.append(f"  - {sh.name}: {text}")
        return "\n".join(out)

    async def _run(self, name: str, inp: dict[str, Any]) -> str:
        d = self.agent
        await d.gate()
        if name == "get_deck":
            return self._describe()
        if name == "set_text":
            i, shape, text = int(inp["slide"]), str(inp["shape"]), str(inp["text"])
            d.status(f"Slide {i + 1}: {shape}")
            await d.set_text(i, shape, text, label=f"slide {i + 1}")
            return "ok"
        if name == "add_slide":
            prs = Presentation(d.pptx)
            n = _add_slide(prs, inp["title"], inp.get("body", ""))
            prs.save(d.pptx)
            await d.render()
            d.emit("deck_loaded", slides=len(prs.slides), base=d.base_url, version=d.version)
            await d._focus(n, [0.06, 0.08, 0.88, 0.17], f"slide {n + 1}")
            d.status(f"Added slide {n + 1}")
            return f"ok: slide {n} added with shapes s{n}_title, s{n}_body"
        if name == "ask_human":
            ans = await d.ask(inp["question"], inp.get("options") or [], timeout=self.answer_timeout)
            self.human_notes.append(f"{d.name} asked: {inp['question']} → {ans}")
            return f"The human answered: {ans}"
        return f"error: unknown tool {name}"

    async def run(self, goal: str, context: str = "") -> WorkerResult:
        d = self.agent
        t0 = time.monotonic()
        messages: list[dict[str, Any]] = [{"role": "user", "content": (
            f"GOAL: {goal}\n\n" + (f"Context from teammates:\n{context}\n\n" if context else "")
            + f"The deck is {d.pptx.name}.")}]
        summary, ok = "", False
        for turn in range(MAX_TURNS):
            await d.gate()
            for m in d.drain_inbox():
                messages.append({"role": "user", "content": f"Message from the human: {m.get('text')}"})
            d.status("Thinking…", min(0.9, 0.05 + turn / MAX_TURNS))
            msg = await llm.client().messages.create(model=llm.WORKER_MODEL, max_tokens=8000,
                                                     system=SYSTEM.format(name=d.name), tools=TOOLS, messages=messages)
            llm.USAGE.add(llm.WORKER_MODEL, msg.usage)
            messages.append({"role": "assistant", "content": msg.content})
            uses = [b for b in msg.content if getattr(b, "type", "") == "tool_use"]
            if not uses:
                summary, ok = llm.text_of(msg).strip()[:400] or "Done.", True
                break
            results, finished = [], False
            for u in uses:
                if u.name == "finish":
                    summary, ok, finished = str(u.input.get("summary", "")), True, True
                    results.append({"type": "tool_result", "tool_use_id": u.id, "content": "ok"})
                    continue
                try:
                    out = await self._run(u.name, dict(u.input))
                except Exception as exc:
                    out = f"error: {exc}"
                results.append({"type": "tool_result", "tool_use_id": u.id, "content": out[:20000]})
            messages.append({"role": "user", "content": results})
            if finished:
                break
        d.status("Done" if ok else "Stopped", 1.0)
        return WorkerResult(d.id, ok, summary or "Stopped.", detail=summary + f"\nDeck: {d.pptx.name}",
                            files=[str(d.pptx)], seconds=time.monotonic() - t0)


async def wait_native_viewer(bus, seconds: float = 2.0) -> bool:
    for _ in range(int(seconds * 10)):
        if bus.deck_viewer:
            return True
        await asyncio.sleep(0.1)
    return bool(bus.deck_viewer)
