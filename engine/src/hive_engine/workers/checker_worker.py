"""Checker worker: reads what the other agents produced (their reports + the files they touched) and
looks for contradictions or mistakes. If a human decision is needed it asks with 2–4 options plus
"or tell me something else" (allow_other)."""

from __future__ import annotations

import time
from pathlib import Path
from typing import Any

from ..orchestrator import llm
from .base import DynamicAgent, WorkerResult
from .window import AgentWindow

SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "ok": {"type": "boolean", "description": "true if everything is consistent and correct"},
        "findings": {"type": "array", "items": {"type": "string"}, "description": "short, specific findings"},
        "question": {"type": "string", "description": "a decision only the human can make (optional)"},
        "options": {"type": "array", "items": {"type": "string"}, "description": "2-4 options, best first"},
        "summary": {"type": "string"},
    },
    "required": ["ok", "findings", "summary"],
}

SYSTEM = """You are {name}, the checker on the Hive team. Compare the teammates' results and files against the goal.
Flag contradictions (e.g. different prices in two places), factual slips and anything left unfinished.
Be concrete and brief. Only ask the human when a real decision is needed; give 2-4 options, best first."""


def _snippets(ws: Path | None, files: list[str], limit: int = 6000) -> str:
    if not ws:
        return ""
    out, used = [], 0
    for rel in files:
        p = (ws / rel) if not Path(rel).is_absolute() else Path(rel)
        if p.suffix == ".pptx":
            try:
                from pptx import Presentation
                text = "\n".join(sh.text_frame.text for s in Presentation(p).slides for sh in s.shapes if sh.has_text_frame)
            except Exception:
                continue
        else:
            try:
                text = p.read_text(errors="replace")
            except OSError:
                continue
        chunk = f"--- {rel} ---\n{text[:2500]}"
        if used + len(chunk) > limit:
            break
        out.append(chunk)
        used += len(chunk)
    return "\n".join(out)


class CheckerWorker:
    def __init__(self, agent: DynamicAgent, ws: Path | None, window: AgentWindow | None, answer_timeout: float = 30) -> None:
        self.agent = agent
        self.ws = ws
        self.win = window
        self.answer_timeout = answer_timeout
        self.human_notes: list[str] = []

    async def run(self, goal: str, team_report: str, files: list[str]) -> WorkerResult:
        a = self.agent
        t0 = time.monotonic()
        a.status("Reviewing the team's work…", 0.2)
        if self.win:
            self.win.bench.status = "review"
            self.win.bench.log("review " + ", ".join(files[:8]), cmd=True)
            await self.win.refresh()
        prompt = (f"TEAM GOAL: {goal}\n\nYOUR TASK: {a.spec.goal}\n\nTEAM RESULTS:\n{team_report}\n\n"
                  f"FILES:\n{_snippets(self.ws, files)}")
        raw = await llm.call_tool(SYSTEM.format(name=a.name), prompt, "submit_review", SCHEMA, "Submit the review.",
                                  model=llm.WORKER_MODEL)
        raw = raw or {"ok": True, "findings": [], "summary": "No review."}
        findings = [str(f) for f in raw.get("findings") or []]
        if self.win:
            for f in findings:
                self.win.bench.log("• " + f)
            await self.win.refresh()
            await self.win.point("#term div:last-of-type", label="finding")
        a.status(f"{len(findings)} finding(s)", 0.7)
        summary = str(raw.get("summary") or ("All consistent." if raw.get("ok") else "Issues found."))
        if raw.get("question") and raw.get("options"):
            ans = await a.ask(str(raw["question"]), [str(o) for o in raw["options"]], timeout=self.answer_timeout)
            self.human_notes.append(f"{a.name} asked: {raw['question']} → {ans}")
            summary += f" Human decided: {ans}."
        a.status("Done", 1.0)
        detail = summary + ("\nFindings:\n- " + "\n- ".join(findings) if findings else "")
        return WorkerResult(a.id, True, summary, detail=detail, seconds=time.monotonic() - t0)
