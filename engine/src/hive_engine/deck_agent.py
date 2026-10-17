"""Deck agent: edits a real .pptx with python-pptx in small steps, re-renders after each one,
and tells the Deck Viewer (and the ghost-cursor overlay) where the edit is happening.

Extra event types used here (not yet in docs/CONTRACT.md — Mac app should ignore unknown types):
  deck_loaded         {"slides": n, "base": url, "version": int}
  deck_focus          {"slide": i, "box": [fx, fy, fw, fh], "label": str}   # fractions of the slide
  deck_slide_updated  {"slide": i, "version": int, "render_ms": int}
"""

from __future__ import annotations

import asyncio
import time
from pathlib import Path
from typing import Callable

from pptx import Presentation
from pptx.util import Emu

from .agent import Agent
from .bus import Bus
from .render import render_pptx

# A function that maps a slide-fraction box to SCREEN points, if the viewer's position is known.
ScreenMapper = Callable[[float, float, float, float], "asyncio.Future | tuple[float, float, float, float]"]


class DeckAgent(Agent):
    def __init__(self, bus: Bus, pptx: Path, render_dir: Path, base_url: str = "http://127.0.0.1:8766/render") -> None:
        super().__init__(bus, "deck")
        self.pptx = pptx
        self.render_dir = render_dir
        self.base_url = base_url
        self.version = 0
        self.screen_mapper = None  # async (fx, fy, fw, fh) -> (x, y, w, h) in screen points
        self.render_times: list[float] = []

    # ---- rendering ---------------------------------------------------------
    async def render(self) -> float:
        t0 = time.perf_counter()
        await asyncio.to_thread(render_pptx, self.pptx, self.render_dir)
        self.version += 1
        dt = time.perf_counter() - t0
        self.render_times.append(dt)
        return dt

    async def load(self) -> None:
        await self.render()
        n = len(Presentation(self.pptx).slides)
        self.emit("deck_loaded", slides=n, base=self.base_url, version=self.version)

    # ---- geometry ----------------------------------------------------------
    def _box_frac(self, prs, shape) -> list[float]:
        sw, sh = prs.slide_width, prs.slide_height
        return [shape.left / sw, shape.top / sh, shape.width / sw, shape.height / sh]

    @staticmethod
    def _find_shape(slide, key):
        if isinstance(key, int):
            return slide.shapes[key]
        for s in slide.shapes:
            if s.name == key:
                return s
            if s.has_text_frame and key.lower() in s.text_frame.text.lower():
                return s  # match by placeholder text, e.g. "TODO: market size"
        raise KeyError(f"shape {key!r} not found")

    async def _focus(self, slide_idx: int, box: list[float], label: str) -> None:
        self.emit("deck_focus", slide=slide_idx, box=[round(b, 4) for b in box], label=label)
        if self.screen_mapper:
            x, y, w, h = await self.screen_mapper(*box)
            await self.move_to(x + w / 2, y + h / 2, 400, label)
            self.highlight(x, y, w, h, style="select")
        else:
            await asyncio.sleep(0.4)

    # ---- steps ---------------------------------------------------------------
    async def set_text(self, slide_idx: int, shape_key, text: str, label: str | None = None) -> None:
        await self.gate()
        prs = Presentation(self.pptx)
        shape = self._find_shape(prs.slides[slide_idx], shape_key)
        await self._focus(slide_idx, self._box_frac(prs, shape), label or "text")
        if self.screen_mapper and self.pos:
            self.emit("typing", x=self.pos[0], y=self.pos[1], text=text)
        shape.text_frame.text = text
        prs.save(self.pptx)
        dt = await self.render()
        self.emit("deck_slide_updated", slide=slide_idx, version=self.version, render_ms=int(dt * 1000))

    async def insert_image(self, slide_idx: int, shape_key, image: Path, label: str | None = None) -> None:
        """Put an image exactly where a placeholder box sits (then remove the box)."""
        await self.gate()
        prs = Presentation(self.pptx)
        slide = prs.slides[slide_idx]
        shape = self._find_shape(slide, shape_key)
        await self._focus(slide_idx, self._box_frac(prs, shape), label or "image")
        if getattr(shape, "is_placeholder", False) and hasattr(shape, "insert_picture"):
            shape.insert_picture(str(image))
        else:
            slide.shapes.add_picture(str(image), Emu(shape.left), Emu(shape.top), Emu(shape.width), Emu(shape.height))
            shape._element.getparent().remove(shape._element)
        prs.save(self.pptx)
        dt = await self.render()
        self.emit("deck_slide_updated", slide=slide_idx, version=self.version, render_ms=int(dt * 1000))
