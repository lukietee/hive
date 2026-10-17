"""deck_sim — the Deck agent fills a sample deck step by step; the stopgap Deck Viewer
(Chrome window) crossfades each render, and ghost-cursor events land on the edited shape."""

from __future__ import annotations

import asyncio
import tempfile
from pathlib import Path

from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from ..browser_agent import Rect, launch_chrome, open_window, page_screen_rect
from ..bus import Bus
from ..deck_agent import DeckAgent
from ..http_server import HTTP_PORT

VIEWER_URL = f"http://127.0.0.1:{HTTP_PORT}/viewer/index.html"


def make_sample_deck(path: Path) -> None:
    """4 slides with TODO placeholders (a stand-in until demo-world/pitch.pptx exists)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for title, body in [("Acme", "TODO: tagline"), ("Problem", "TODO: problem"),
                        ("Market", "TODO: market size"), ("Pricing", "TODO: pricing")]:
        s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        tb = s.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(1.4))
        tb.name = "title"
        tb.text_frame.text = title
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(48)
        tb.text_frame.paragraphs[0].runs[0].font.bold = True
        bb = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(3.5))
        bb.name = "body"
        bb.text_frame.text = body
        bb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        bb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    prs.save(path)


async def run(bus: Bus, render_dir: Path, pptx: Path | None = None, keep_open: float = 2.0, **_) -> dict:
    if pptx is None:
        pptx = Path(tempfile.gettempdir()) / "hive_sample_deck.pptx"
        make_sample_deck(pptx)
    deck = DeckAgent(bus, pptx, render_dir)

    async with async_playwright() as pw:
        browser = await launch_chrome(pw)
        ctx, page = await open_window(browser, Rect(80, 80, 1100, 700), VIEWER_URL)

        async def mapper(fx: float, fy: float, fw: float, fh: float):
            x, y, w, h = await page_screen_rect(page, page.locator("#stage"))
            return x + fx * w, y + fy * h, fw * w, fh * h

        deck.screen_mapper = mapper
        await asyncio.sleep(1.0)  # let the viewer connect to the bus
        x, y, w, h = await page_screen_rect(page, page.locator("#stage"))
        deck.spawn(x + 30, y + 30)
        await deck.load()
        deck.status("Filling blank slides", 0.1)
        steps = [(0, "body", "AI agents that finish your launch while you pitch"),
                 (1, "body", "Launches need 5 deliverables at once; tools do one at a time"),
                 (2, "body", "$4B — teams shipping weekly"),
                 (3, "body", "Free · Pro $29/mo · Team $99/mo")]
        for i, (slide, shape, text) in enumerate(steps):
            await deck.set_text(slide, shape, text, label=f"slide {slide + 1}")
            deck.status(f"Slide {slide + 1} done", (i + 1) / len(steps))
        await asyncio.sleep(0.6)
        await page.screenshot(path=str(render_dir / "viewer_check.png"))  # debug: what the viewer showed
        deck.done("4 slides filled")
        await asyncio.sleep(keep_open)
        await ctx.close()
        await browser.close()
    t = deck.render_times
    return {"renders": len(t), "first_s": round(t[0], 2), "avg_s": round(sum(t[1:]) / max(1, len(t) - 1), 2), "pptx": str(pptx)}
